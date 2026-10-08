"""Idempotently add security cases to the existing workbook without erasing execution history."""
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

path = Path(__file__).resolve().parents[1] / 'test-case' / 'Login_TestCases.xlsx'
wb = load_workbook(path)
tc, ex = wb['TestCases'], wb['Execution']

cases = [
    ('SEC-LOGIN-001', 'SQL injection', 'Dấu nháy đơn trong username', 'Phát hiện lỗi SQL lộ ra giao diện hoặc vượt qua xác thực', 'Cần staging được phép kiểm thử; bật UTC_SECURITY_TESTS_ENABLED và UTC_SECURITY_URL.', "'", 'TestOnly!2026', '1. Mở Login staging\n2. Nhập dấu nháy đơn vào username, mật khẩu giả\n3. Gửi form một lần\n4. Kiểm tra lỗi và form', 'Form còn hiển thị; có lỗi từ ứng dụng; không lộ lỗi SQL/DB.', 'OWASP WSTG; giả định cần xác nhận trên staging, chưa chạy production.', 'Critical', 'Có điều kiện: staging', 'UsernameQuoteDoesNotExposeSqlErrorOrAuthenticate'),
    ('SEC-LOGIN-002', 'SQL injection', 'Dấu nháy đơn trong password', 'Phát hiện lỗi SQL lộ ra giao diện hoặc vượt qua xác thực', 'Cần staging được phép kiểm thử; bật UTC_SECURITY_TESTS_ENABLED và UTC_SECURITY_URL.', 'utc_qa_nonexistent_20261008', "'", '1. Mở Login staging\n2. Nhập username giả, dấu nháy đơn vào password\n3. Gửi form một lần\n4. Kiểm tra lỗi và form', 'Form còn hiển thị; có lỗi từ ứng dụng; không lộ lỗi SQL/DB.', 'OWASP WSTG; giả định cần xác nhận trên staging, chưa chạy production.', 'Critical', 'Có điều kiện: staging', 'PasswordQuoteDoesNotExposeSqlErrorOrAuthenticate'),
    ('SEC-LOGIN-003', 'Rate limiting', 'Giới hạn thử đăng nhập liên tiếp', 'Kiểm tra tín hiệu giới hạn theo ngưỡng được phê duyệt', 'Cần staging, tài khoản thử, ngưỡng và tín hiệu rate limiting đã xác nhận.', 'Tài khoản thử trên staging', 'Mật khẩu giả', '1. Phê duyệt ngưỡng và số lần gửi\n2. Chạy tuần tự trên staging\n3. Dừng khi có CAPTCHA/khóa/rate limit\n4. So sánh với ngưỡng', 'Tín hiệu rate limit/CAPTCHA xuất hiện theo ngưỡng đã thống nhất.', 'OWASP WSTG; ngưỡng và kỳ vọng chưa được xác nhận.', 'Critical', 'Blocked: cần staging và ngưỡng', ''),
    ('SEC-LOGIN-004', 'DDoS resilience', 'Duy trì dịch vụ trong tải đã phê duyệt', 'Đánh giá độ sẵn sàng với load profile và SLA được phê duyệt', 'Cần môi trường cô lập, load profile, SLA và phương án dừng đã phê duyệt.', 'Không dùng tài khoản thật', 'Không áp dụng', '1. Phê duyệt kế hoạch tải\n2. Chạy bằng công cụ performance\n3. Theo dõi latency, error rate, tài nguyên\n4. Dừng tại ngưỡng an toàn và so sánh SLA', 'Đạt SLA đã thống nhất; không ảnh hưởng người dùng khác.', 'Kỳ vọng cần xác nhận; chưa có môi trường và SLA.', 'Normal', 'Blocked: cần staging, load plan và SLA', ''),
]

existing = {tc.cell(row, 1).value for row in range(2, tc.max_row + 1)}
new = [row for row in cases if row[0] not in existing]
if not new:
    print('Security cases already present; workbook unchanged.')
    raise SystemExit(0)

latest = next((ex.cell(row, 1).value for row in range(ex.max_row, 1, -1)
               if ex.cell(row, 1).value not in (None, 'INITIAL')), None)

for values in new:
    tc.append(values)
    row = tc.max_row
    tc.row_dimensions[row].height = 94
    for col in range(1, 14):
        tc.cell(row, col)._style = copy(tc.cell(2, col)._style)
tc.auto_filter.ref = f'A1:M{tc.max_row}'

# Keep INITIAL before historical runs so Summary still points to the latest real Run ID.
initial_insert_at = 2
while initial_insert_at <= ex.max_row and ex.cell(initial_insert_at, 1).value == 'INITIAL':
    initial_insert_at += 1
ex.insert_rows(initial_insert_at, amount=len(new))
for offset, values in enumerate(new):
    row = initial_insert_at + offset
    data = ['INITIAL', values[0], '', '', 'Chưa chạy project NUnit', 'Not Run',
            'Chưa có staging/ngưỡng/SLA được xác nhận; chưa thực thi.', '', '']
    ex.row_dimensions[row].height = 54
    for col, value in enumerate(data, 1):
        ex.cell(row, col).value = value
        ex.cell(row, col)._style = copy(ex.cell(2, col)._style)

if latest:
    for values in new:
        ex.append([latest, values[0], '', '', 'Case được thêm sau lần chạy này.', 'Not Run',
                   'Chưa tồn tại khi Run ID này được thực thi; chưa chạy.', '', ''])
        row = ex.max_row
        ex.row_dimensions[row].height = 54
        for col in range(1, 10):
            ex.cell(row, col)._style = copy(ex.cell(2, col)._style)

ex.auto_filter.ref = f'A1:I{max(ex.max_row, 5000)}'
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(path)
print(f'Added {len(new)} security cases to {path}; latest historical run: {latest or "none"}.')

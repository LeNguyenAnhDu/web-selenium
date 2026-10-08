from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from pathlib import Path

root = Path(__file__).resolve().parents[1]
path = root / 'test-case' / 'Login_TestCases.xlsx'
path.parent.mkdir(exist_ok=True)

cases = [
    ('TC-LOGIN-001', 'Validation server', 'Bỏ trống cả hai trường', 'Xác nhận lỗi khi không nhập username/password', '', '', 'Bạn chưa nhập tên đăng nhập', 'Critical', 'Có', 'BothFieldsEmptyShowsServerValidation'),
    ('TC-LOGIN-002', 'Validation server', 'Có username giả, bỏ trống mật khẩu', 'Xác nhận lỗi thiếu mật khẩu', 'utc_qa_nonexistent_20261008', '', 'Bạn chưa nhập mật khẩu', 'Critical', 'Có', 'PasswordEmptyShowsServerValidation'),
    ('TC-LOGIN-003', 'Validation server', 'Bỏ trống username, có mật khẩu giả', 'Xác nhận lỗi thiếu username', '', 'TestOnly!2026', 'Bạn chưa nhập tên đăng nhập', 'Critical', 'Có', 'UsernameEmptyShowsServerValidation'),
    ('TC-LOGIN-004', 'Xác thực thất bại', 'Tài khoản giả dùng riêng cho kiểm thử', 'Xác nhận phản hồi xác thực thất bại', 'utc_qa_nonexistent_20261008', 'TestOnly!2026', 'Tài khoản hoặc mật khẩu không đúng.', 'Critical', 'Có', 'FakeAccountShowsAuthenticationError'),
    ('TC-LOGIN-005', 'Xác thực thất bại', 'Hai trường chỉ chứa khoảng trắng', 'Xác nhận hành vi thực tế với khoảng trắng', '[3 dấu cách]', '[3 dấu cách]', 'Tài khoản hoặc mật khẩu không đúng.', 'Normal', 'Có', 'WhitespaceFieldsShowAuthenticationError'),
    ('TC-LOGIN-006', 'Giao diện', 'Mật khẩu được che', 'Xác nhận ô mật khẩu che ký tự khi nhập', '', 'TestOnly!2026', 'Ô mật khẩu có type=password sau khi nhập; không gửi form.', 'Normal', 'Có', 'PasswordInputIsMasked'),
]

wb = Workbook()
tc = wb.active
tc.title = 'TestCases'
ex = wb.create_sheet('Execution')
su = wb.create_sheet('Summary')
tc.append(['Test Case ID', 'Nhóm kiểm thử', 'Tên test case', 'Mục tiêu', 'Điều kiện tiên quyết', 'Dữ liệu username', 'Dữ liệu password giả', 'Các bước thực hiện', 'Kết quả mong đợi', 'Cơ sở xác định kết quả mong đợi', 'Mức ưu tiên', 'Khả năng automation', 'Tên phương thức test tương ứng'])
ex.append(['Run ID', 'Test Case ID', 'Thời điểm chạy và múi giờ', 'Môi trường/trình duyệt', 'Kết quả thực tế', 'Trạng thái', 'Nguyên nhân lỗi hoặc lý do chưa chạy', 'Đường dẫn screenshot/evidence', 'Đường dẫn hoặc mã tham chiếu kết quả Allure'])

for id, group, name, goal, username, password, expected, priority, automation, method in cases:
    steps = '1. Mở https://vanphongdientu.utc.edu.vn/Login\n2. Nhập dữ liệu trong cột username/password\n'
    steps += '3. Gửi form\n4. Chờ phản hồi\n5. Kiểm tra lỗi và form' if id != 'TC-LOGIN-006' else '3. Kiểm tra kiểu ô mật khẩu; không gửi form'
    result = expected if id == 'TC-LOGIN-006' else f'Hiển thị div.error: “{expected}”; form đăng nhập vẫn hiển thị.'
    basis = 'DOM và hành vi quan sát bằng Chrome headless ngày 08/10/2026; không phải yêu cầu sản phẩm đã xác nhận.'
    tc.append([id, group, name, goal, 'Truy cập được trang; không có CAPTCHA/khóa/rate limit.', username, password, steps, result, basis, priority, automation, method])
    ex.append(['INITIAL', id, '', '', 'Chưa chạy project NUnit', 'Not Run', 'Chưa thực thi project kiểm thử', '', ''])

su['A1'] = 'TỔNG HỢP KIỂM THỬ ĐĂNG NHẬP UTC'
su['A3'] = 'Tổng số test case'
su['B3'] = '=COUNTA(TestCases!A2:A5000)'
su['A5'] = 'Mức ưu tiên'
su['B5'] = 'Số lượng'
for row, priority in [(6, 'Critical'), (7, 'Normal')]:
    su[f'A{row}'] = priority
    su[f'B{row}'] = f'=COUNTIF(TestCases!K2:K5000,A{row})'
su['A9'] = 'Run ID gần nhất'
su['B9'] = '=IFERROR(LOOKUP(2,1/(Execution!A2:A5000<>""),Execution!A2:A5000),"")'
su['A11'] = 'Trạng thái'
su['B11'] = 'Số lượng'
for row, status in enumerate(['Not Run', 'Passed', 'Failed', 'Blocked', 'Skipped'], 12):
    su[f'A{row}'] = status
    su[f'B{row}'] = f'=COUNTIFS(Execution!A2:A5000,$B$9,Execution!F2:F5000,A{row})'
su['A18'] = 'Chỉ tính các dòng của Run ID gần nhất. Dòng INITIAL lưu trạng thái ban đầu.'

header_fill = PatternFill('solid', fgColor='17365D')
white = Font(color='FFFFFF', bold=True)
thin = Side(style='thin', color='D9E2F3')
for ws, widths in [(tc, [18, 22, 34, 38, 42, 30, 22, 55, 55, 65, 16, 21, 38]),
                   (ex, [24, 18, 29, 34, 55, 15, 55, 62, 62])]:
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = white
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 36
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 94 if ws == tc else 42
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = Border(bottom=thin)
    ws.sheet_view.showGridLines = False
ex.auto_filter.ref = 'A1:I5000'
validation = DataValidation(type='list', formula1='"Not Run,Passed,Failed,Blocked,Skipped"')
ex.add_data_validation(validation)
validation.add('F2:F5000')
colors = {'Not Run': 'E7E6E6', 'Passed': 'C6EFCE', 'Failed': 'FFC7CE', 'Blocked': 'FFE699', 'Skipped': 'D9EAD3'}
for status, color in colors.items():
    ex.conditional_formatting.add('F2:F5000', CellIsRule(operator='equal', formula=[f'"{status}"'], fill=PatternFill('solid', fgColor=color)))
su.column_dimensions['A'].width = 70
su.column_dimensions['B'].width = 24
su.merge_cells('A1:B1')
su['A1'].fill = header_fill
su['A1'].font = Font(color='FFFFFF', bold=True, size=15)
su.row_dimensions[1].height = 32
for row in (5, 11):
    for cell in su[row][:2]:
        cell.fill = header_fill
        cell.font = white
su.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(path)
print(path)

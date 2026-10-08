"""Append one real NUnit run to Login_TestCases.xlsx; never infer Passed from missing data."""
import json
import re
import sys
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta, timezone
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

run_dir = Path(sys.argv[1]).resolve()
run_id = sys.argv[2]
root = Path(__file__).resolve().parents[1]
book_path = root / 'test-case' / 'Login_TestCases.xlsx'
trx_path = run_dir / 'results.trx'
if not trx_path.exists():
    raise SystemExit(f'Không tìm thấy TRX: {trx_path}; workbook chưa cập nhật.')

tree = ET.parse(trx_path)
ns = {'t': 'http://microsoft.com/schemas/VisualStudio/TeamTest/2010'}
results = {}
method_to_id = {
    'FakeAccountShowsAuthenticationError': 'TC-LOGIN-004',
    'WhitespaceFieldsShowAuthenticationError': 'TC-LOGIN-005',
    'PasswordInputIsMasked': 'TC-LOGIN-006',
}
for item in tree.findall('.//t:UnitTestResult', ns):
    name = item.get('testName', '')
    match = re.search(r'TC-LOGIN-\d{3}', name)
    id = match.group() if match else method_to_id.get(name)
    if not id:
        continue
    message = item.find('.//t:Message', ns)
    results[id] = (item.get('outcome', ''), message.text.strip() if message is not None and message.text else '', item.get('startTime', ''))

allure = {}
for path in (run_dir / 'allure-results').glob('*-result.json'):
    data = json.loads(path.read_text(encoding='utf-8-sig'))
    labels = data.get('labels', [])
    ids = [x.get('value') for x in labels if x.get('name') == 'testCaseId']
    if ids:
        allure[ids[0]] = str(path.relative_to(root))

wb = load_workbook(book_path)
ws = wb['Execution']
existing_rows = {ws.cell(row, 2).value: row for row in range(2, ws.max_row + 1) if ws.cell(row, 1).value == run_id}
ids = [wb['TestCases'].cell(row, 1).value for row in range(2, wb['TestCases'].max_row + 1)]
now = datetime.now(timezone(timedelta(hours=7))).strftime('%d/%m/%Y %H:%M:%S UTC+07:00')
environment = 'Windows / Chrome / ' + ('Headless' if __import__('os').environ.get('UTC_HEADLESS', 'true') == 'true' else 'Có giao diện')

for id in ids:
    outcome, message, start_time = results.get(id, ('', '', ''))
    if not outcome:
        status, actual, reason = 'Not Run', 'Không có UnitTestResult tương ứng trong TRX.', 'Không được thực thi.'
    elif 'Blocked:' in message:
        status, actual, reason = 'Blocked', 'Kiểm thử dừng trước assertion nghiệp vụ.', message
    elif outcome == 'Passed':
        status, actual, reason = 'Passed', 'Các assertion NUnit của test case đã đạt.', ''
    elif outcome == 'Failed':
        status, actual, reason = 'Failed', 'Assertion hoặc thao tác Selenium thất bại.', message
    else:
        status, actual, reason = 'Skipped', f'NUnit outcome: {outcome}', message or 'NUnit không chạy test.'
    screenshots = sorted((run_dir / 'screenshots').glob(f'{id}_*.png'))
    evidence = str(screenshots[-1].relative_to(root)) if screenshots else ''
    run_time = start_time.replace('T', ' ') if start_time else now
    values = [run_id, id, run_time, environment, actual, status, reason, evidence, allure.get(id, '')]
    row = existing_rows.get(id)
    if row:
        for column, value in enumerate(values, 1):
            ws.cell(row, column, value)
    else:
        ws.append(values)
        row = ws.max_row
    ws.row_dimensions[row].height = 54
    for cell in ws[row]:
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        cell.border = Border(bottom=Side(style='thin', color='D9E2F3'))

ws.auto_filter.ref = f'A1:I{ws.max_row}'
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(book_path)
print(f'Workbook updated: {book_path}; {len(results)} NUnit results; {len(allure)} Allure results')
